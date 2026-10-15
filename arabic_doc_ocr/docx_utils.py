"""DOCX helpers: right-to-left paragraphs, tables from the HTML / Markdown that dots.mocr produces."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

_ARABIC = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
_LATIN = re.compile(r"[A-Za-z]")
_HTML_TABLE = re.compile(r"<table\b.*?</table>", re.S | re.I)
_TAG = re.compile(r"</?[a-zA-Z][a-zA-Z0-9]*(?:\s[^<>]*)?/?>")
_MD_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?$")


@dataclass
class Block:
    kind: str  # "heading" | "para" | "table"
    text: str = ""
    level: int = 0
    rows: list | None = None  # table: list[list[(text, colspan)]]


_STRONG = re.compile(r"[A-Za-z\u00C0-\u024F\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")


def para_rtl(text: str) -> bool:
    """Paragraph direction per UAX #9 rule P2: the first strong character is Arabic / Hebrew → right-to-left."""
    m = _STRONG.search(text or "")
    return bool(m) and m.group(0) >= "\u0590"


def is_rtl(text: str) -> bool:
    return len(_ARABIC.findall(text)) > len(_LATIN.findall(text))


def _list_marker(kind: str, n: int) -> str:
    """List marker of an <li> by <ol type>: 1 → 1. · a → a. · A → A. · i → i. · I → I. · ul → •"""
    if kind == "ul":
        return "•"
    if kind in ("a", "A"):
        s, x = "", n
        while x > 0:
            x, r = divmod(x - 1, 26)
            s = chr(97 + r) + s
        return (s if kind == "a" else s.upper()) + "."
    if kind in ("i", "I"):
        vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
                (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
        s, x = "", n
        for v, sym in vals:
            while x >= v:
                s, x = s + sym, x - v
        return (s if kind == "i" else s.upper()) + "."
    return f"{n}."


def _cell_text(el) -> str:
    """Text of an HTML table cell, keeping line breaks (<br>, <p>) and list numbering (nested lists indented)."""
    out: list[str] = []

    def nl():
        if out and not out[-1].endswith("\n"):
            out.append("\n")

    def walk(node, depth):
        tag = node.tag.lower() if isinstance(node.tag, str) else ""
        if tag == "br":
            nl()
        elif tag in ("ol", "ul"):
            kind = "ul" if tag == "ul" else (node.get("type") or "1")
            try:
                n = int(node.get("start") or 1)
            except ValueError:
                n = 1
            if node.text and node.text.strip():
                out.append(node.text)
            for li in node:
                if not isinstance(li.tag, str) or li.tag.lower() != "li":
                    continue
                nl()
                out.append("    " * depth + _list_marker(kind, n) + " ")
                if li.text:
                    out.append(li.text)
                for ch in li:
                    walk(ch, depth + 1)
                n += 1
            nl()
        else:
            block = tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6")
            if block:
                nl()
            if node.text:
                out.append(node.text)
            for ch in node:
                walk(ch, depth)
            if block:
                nl()
        if node.tail:
            out.append(node.tail)

    if el.text:
        out.append(el.text)
    for ch in el:
        walk(ch, 0)
    lines = []
    for line in "".join(out).split("\n"):
        indent = len(line) - len(line.lstrip(" "))
        body = " ".join(line.split())
        if body:
            lines.append(" " * indent + body)
    return "\n".join(lines)


def _html_table_rows(table_html: str) -> list[list[tuple[str, int]]]:
    from lxml import html as lh

    doc = lh.document_fromstring(f"<html><body>{table_html}</body></html>")
    rows = []
    for tr in doc.body.xpath(".//tr"):
        cells = []
        for td in tr.xpath("./td|./th"):
            try:
                span = max(1, int(td.get("colspan", 1)))
            except ValueError:
                span = 1
            cells.append((_cell_text(td), span))
        if cells:
            rows.append(cells)
    return rows


def _text_blocks(segment: str) -> list[Block]:
    """Non-HTML-table part: headings, Markdown tables, paragraphs."""
    segment = re.sub(r"<br\s*/?>", "\n", segment, flags=re.I)
    segment = re.sub(r"</(?:p|div|h[1-6]|li)>", "\n\n", segment, flags=re.I)
    segment = html.unescape(_TAG.sub("", segment))
    blocks, para, md_rows = [], [], []

    def flush_para():
        if para:
            blocks.append(Block("para", "\n".join(para)))
            para.clear()

    def flush_table():
        if md_rows:
            blocks.append(Block("table", rows=[[(c.strip(), 1) for c in r] for r in md_rows]))
            md_rows.clear()

    for raw in segment.splitlines():
        line = raw.strip()
        if line.startswith("```"):
            continue
        if line.startswith("|") and line.count("|") >= 2:
            flush_para()
            if not _MD_SEP.match(line):
                md_rows.append(line.strip("|").split("|"))
            continue
        flush_table()
        if not line:
            flush_para()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_para()
            blocks.append(Block("heading", m.group(2).strip(), level=min(len(m.group(1)), 3)))
            continue
        para.append(re.sub(r"^[*+]\s+", "• ", re.sub(r"^-\s+", "• ", line)))
    flush_para()
    flush_table()
    return blocks


def parse_blocks(text: str) -> list[Block]:
    blocks, pos = [], 0
    for m in _HTML_TABLE.finditer(text):
        blocks += _text_blocks(text[pos:m.start()])
        rows = _html_table_rows(m.group(0))
        if rows:
            blocks.append(Block("table", rows=rows))
        pos = m.end()
    blocks += _text_blocks(text[pos:])
    return blocks


# --- writing ---------------------------------------------------------------------------


# CT_PPrBase: elements that come AFTER w:bidi (wrong order → Word reports the file as damaged)
_AFTER_BIDI = ("adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
               "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
               "cnfStyle", "rPr", "sectPr", "pPrChange")


def _set_bidi(paragraph) -> None:
    """Right-to-left paragraph aligned to its start (= right). jc="start", not "right": in a bidi paragraph
    left/right are logical, so "right" would align left in Word / LibreOffice."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # creates w:jc at the right position, value changed below
    ppr = paragraph._p.get_or_add_pPr()
    ppr.find(qn("w:jc")).set(qn("w:val"), "start")
    if ppr.find(qn("w:bidi")) is None:
        el = OxmlElement("w:bidi")
        nxt = next((c for c in ppr if isinstance(c.tag, str) and c.tag.split("}")[-1] in _AFTER_BIDI), None)
        if nxt is not None:
            nxt.addprevious(el)
        else:
            ppr.append(el)


def _add_runs(paragraph, text: str) -> None:
    """**bold** → bold run; newline → soft line break."""
    for li, line in enumerate(text.split("\n")):
        if li:
            paragraph.add_run().add_break()
        for part in re.split(r"(\*\*[^*]+\*\*)", line):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                paragraph.add_run(part[2:-2]).bold = True
            else:
                paragraph.add_run(part.replace("**", ""))


def _add_table(doc, rows) -> None:
    n_cols = max(sum(span for _, span in r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    rtl = is_rtl(" ".join(t for r in rows for t, _ in r))
    if rtl:  # CT_TblPr: bidiVisual after tblStyle / tblpPr / tblOverlap, before tblW
        tblpr = table._tbl.tblPr
        el = OxmlElement("w:bidiVisual")
        nxt = next((c for c in tblpr if isinstance(c.tag, str) and c.tag.split("}")[-1] not in
                    ("tblStyle", "tblpPr", "tblOverlap")), None)
        if nxt is not None:
            nxt.addprevious(el)
        else:
            tblpr.append(el)
    for i, r in enumerate(rows):
        j = 0
        for text, span in r:
            if j >= n_cols:
                break
            cell = table.cell(i, j)
            if span > 1 and j + span - 1 < n_cols:
                cell = cell.merge(table.cell(i, j + span - 1))
            p = cell.paragraphs[0]
            _add_runs(p, text)
            if i == 0 and all(not ch.isdigit() for ch in text):
                for run in p.runs:
                    run.bold = True
            if para_rtl(text):
                _set_bidi(p)
            j += span
    doc.add_paragraph()


def _setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Arial")  # w:cs = font for Arabic (complex script)


def _add_markdown(doc, text: str) -> int:
    n = 0
    for b in parse_blocks(text):
        n += 1
        if b.kind == "table":
            _add_table(doc, b.rows)
            continue
        p = doc.add_heading(level=b.level) if b.kind == "heading" else doc.add_paragraph()
        _add_runs(p, b.text)
        if para_rtl(b.text):
            _set_bidi(p)
    return n


def new_document(title: str | None = None):
    doc = Document()
    _setup_styles(doc)
    if title:
        doc.core_properties.title = title
    return doc
