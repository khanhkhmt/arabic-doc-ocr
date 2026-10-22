"""Output files of a document:

    <name>_all.zip          everything below
    <name>.docx             main text, words the cross-check reads differently highlighted yellow + Word comment;
                            with the LLM step: words the LLM changed highlighted green + comment "OCR read: …"
    <name>_main.txt         main text (handwritten blocks: ft3, printed blocks: dots, Hebrew blocks: dots; after the
                            LLM step if it ran)
    <name>_ocr.txt          only with the LLM step: the main text BEFORE the LLM correction
    <name>_ft3.txt / _ft4.txt   each line reader on its own
    <name>_review.html      line-by-line review table (line image | ft3 with highlights | ft4)
    <name>.json             blocks / lines with boxes and every reading
    <name>_pN_boxes.jpg     page with dots blocks and lines drawn (green = ft3 and ft4 agree, orange = differ)
"""

from __future__ import annotations

import base64
import html
import io
import json
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw

from . import config
from .docx_utils import _add_markdown, _set_bidi, new_document
from .pipeline import PageResult
from .text import agreement, compare_words

CSS = """
body{font-family:sans-serif;margin:16px}
.htr table{border-collapse:collapse;width:100%} .htr td,.htr th{border:1px solid #ccc;padding:4px;vertical-align:top}
.htr .ar{direction:rtl;text-align:right;font-family:'Noto Naskh Arabic','Amiri',serif;font-size:17px;line-height:1.7;white-space:pre-wrap}
.htr th{background:#f3f3f3} .htr summary{cursor:pointer;color:#555} .htr img{max-width:100%;border:1px solid #eee}
.htr .meta{color:#666;font-size:12px} .htr mark{background:#ffe14d;padding:0 1px} .htr .ok{color:#1a7f37}
.htr .diff{color:#c2410c} .htr .sum{display:flex;gap:18px;flex-wrap:wrap;margin:6px 0} .htr .sum b{font-size:20px}
.htr tr.block td{background:#eef4ff} .htr mark.llm{background:#86efac}
"""


def _b64(im: Image.Image, h: int = 56) -> str:
    im = im.copy()
    im.thumbnail((900, h))
    buf = io.BytesIO()
    im.convert("RGB").save(buf, format="JPEG", quality=80)
    return base64.b64encode(buf.getvalue()).decode()


def marked_html(main: str, check: str, check_name: str = config.CHECK) -> str:
    return " ".join(f"<mark title='{check_name}: {html.escape(alt)}'>{html.escape(w)}</mark>" if s else html.escape(w)
                    for w, s, alt in compare_words(main, check))


def _llm_rows(res: PageResult, bi: int) -> list[str]:
    """Review-table row with the block text after the LLM step (green = changed, hover = OCR reading)."""
    if not res.llm_changes.get(bi) or "ocr" not in res.texts:
        return []
    main, orig = res.texts["main"].get(bi, ""), res.texts["ocr"].get(bi, "")
    marked = " ".join(f"<mark class=llm title='OCR: {html.escape(alt)}'>{html.escape(w)}</mark>" if s else html.escape(w)
                      for w, s, alt in compare_words(main, orig))
    return [f"<tr><td colspan=3><span class=meta>After the LLM correction — {len(res.llm_changes[bi])} change(s), "
            f"green = changed (hover: what the OCR read)</span><div class=ar>{marked}</div></td></tr>"]


# --- page image with boxes ------------------------------------------------------------------------------------------

def draw_boxes(res: PageResult) -> Image.Image:
    im = res.image.copy()
    dr = ImageDraw.Draw(im)
    for bi, b in enumerate(res.blocks):
        bb = b.get("bbox")
        if bb and len(bb) == 4:
            col = (147, 51, 234) if b.get("hebrew") else (30, 90, 220)
            dr.rectangle([int(round(v)) for v in bb], outline=col, width=3)
            dr.text((bb[0] + 4, bb[1] + 2), str(bi + 1), fill=col)
    ft3, ft4 = res.readings[config.MAIN], res.readings[config.CHECK]
    for n, L in enumerate(res.lines):
        col = (26, 127, 55) if agreement(ft3[n], ft4[n]) >= 0.999 else (234, 88, 12)
        if L.poly:
            dr.polygon([tuple(p) for p in L.poly], outline=col, width=2)
        else:
            dr.rectangle(L.box, outline=col, width=2)
    return im


# --- review HTML ----------------------------------------------------------------------------------------------------

def review_html(res: PageResult, page_no: int, n_pages: int) -> tuple[str, dict]:
    ft3, ft4, lines = res.readings[config.MAIN], res.readings[config.CHECK], res.lines
    words = flagged = lines_differ = 0
    rows = []
    for bi, b in enumerate(res.blocks):
        cat = html.escape(str(b.get("category")))
        idx = [n for n, L in enumerate(lines) if L.block == bi]
        if bi not in res.sources:
            rows.append(f"<tr class=block><td colspan=3>Block {bi + 1} · {cat} — not text (kept as an image in the "
                        "DOCX)</td></tr>")
            continue
        src, agree = res.sources[bi]
        if src == "hebrew":
            rows.append(f"<tr class=block><td colspan=3>Block {bi + 1} · HEBREW — the line readers read Arabic only, "
                        "so this block keeps the dots text (not cross-checked):"
                        f"<div class=ar>{html.escape(b.get('text') or '')}</div></td></tr>")
            continue
        ag = f" · dots ≈ ft3: {100 * agree:.0f}% of words" if agree is not None else ""
        line_rows = []
        for n in idx:
            cw = compare_words(ft3[n], ft4[n])
            m = sum(s for _, s, _ in cw)
            if src == config.MAIN:
                words, flagged, lines_differ = words + len(cw), flagged + m, lines_differ + bool(m)
            tag = "<span class=ok>✓ agree</span>" if not m else f"<span class=diff>{m}/{len(cw)} words differ</span>"
            line_rows.append(f"<tr><td><img src='data:image/jpeg;base64,{_b64(lines[n].image)}'><br>"
                             f"<span class=meta>line {n + 1} · {tag}</span></td>"
                             f"<td class=ar>{marked_html(ft3[n], ft4[n])}</td><td class=ar>{html.escape(ft4[n])}</td></tr>")
        if src == "dots":
            cw = compare_words(b.get("text") or "", " ".join(ft3[n] for n in idx))
            words, flagged = words + len(cw), flagged + sum(s for _, s, _ in cw)
            rows.append(f"<tr class=block><td colspan=3>Block {bi + 1} · {cat} · PRINTED → dots text{ag} · "
                        f"yellow = ft3 reads differently</td></tr><tr><td colspan=3 class=ar>"
                        f"{marked_html(b.get('text') or '', ' '.join(ft3[n] for n in idx), 'ft3')}</td></tr>")
            if idx:
                rows.append(f"<tr><td colspan=3><details><summary>{len(idx)} ft3 / ft4 lines of this block</summary>"
                            f"<table>{''.join(line_rows)}</table></details></td></tr>")
            rows += _llm_rows(res, bi)
            continue
        dots_txt = (f"<details><summary>dots reading of this block (printed-text reader)</summary><div class=ar>"
                    f"{html.escape(b.get('text') or '')}</div></details>") if b.get("text") else ""
        rows.append(f"<tr class=block><td colspan=3>Block {bi + 1} · {cat} · {len(idx)} lines → ft3{ag}"
                    f"{' · FALLBACK (dots found no block)' if b.get('fallback') else ''}{dots_txt}</td></tr>")
        rows += line_rows
        rows += _llm_rows(res, bi)
    count = {x: sum(v[0] == x for v in res.sources.values()) for x in (config.MAIN, "dots", "hebrew")}
    stats = {"lines": len(lines), "lines_differ": lines_differ, "words": words, "words_flagged": flagged,
             "blocks_by_reader": count}
    notes = "".join(f"<p class=meta>⚠ {html.escape(x)}</p>" for x in res.info.get("notes", []))
    i = res.info
    li = i.get("llm") or {}
    if li.get("ok") and li.get("model"):
        n_llm = sum(len(v) for v in res.llm_changes.values())
        notes += (f"<p class=meta>LLM correction ({html.escape(str(li['model']))}, {html.escape(str(li['method']))}): "
                  f"{n_llm} word change(s) on this page — rows \"After the LLM correction\" (green). The line table "
                  "shows the OCR readings before the correction.</p>")
    summary = (f"<div class=sum><span><b>{len(lines)}</b> lines</span><span><b>{lines_differ}</b> lines ft3 ≠ ft4</span>"
               f"<span><b>{100 * flagged / max(1, words):.0f}%</b> words highlighted ({flagged}/{words})</span>"
               f"<span>blocks: <b>{count[config.MAIN]}</b> ft3 · <b>{count['dots']}</b> dots · "
               f"<b>{count['hebrew']}</b> Hebrew</span></div>"
               f"<p class=meta>Page {page_no}/{n_pages} · {len(res.blocks)} blocks · segmentation "
               f"{', '.join(i.get('segmentation') or []) or '—'} · dots ×{i.get('dots_scale', 1):g}, "
               f"{i.get('dots_seconds', 0):.0f} s · ft3 + ft4 {i.get('htr_seconds', 0):.0f} s</p>{notes}"
               "<p class=meta>Yellow = the cross-check reads that word differently (hover to see). On 500-line test "
               "sets ft3 + ft4 highlight ~13–19% of words and catch 57–75% of wrong words — words that are NOT "
               "highlighted can still be wrong, above all digits, names and dates.</p>")
    head = "<tr><th style='width:34%'>line image</th><th>ft3 (main)</th><th>ft4 (check)</th></tr>"
    return f"<div class=htr><h3>Page {page_no}</h3>{summary}<table>{head}{''.join(rows)}</table></div>", stats


# --- DOCX -----------------------------------------------------------------------------------------------------------

def _rtl_run(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))


def _marked_paragraph(doc, main: str, check: str, check_name: str, orig: str | None = None) -> tuple[int, int, int]:
    """One RTL paragraph of the main reading; differing words yellow, each differing span gets one Word comment
    "<check_name> reads: …". orig = the OCR text before the LLM step: words the LLM changed are green, with one
    comment "LLM changed — OCR read: …" per changed span (green wins over yellow). → (words, yellow, green)."""
    from docx.enum.text import WD_COLOR_INDEX

    p = doc.add_paragraph()
    _set_bidi(p)
    words = compare_words(main, check)
    llm = compare_words(main, orig) if orig is not None else [(w, False, "") for w, _, _ in words]
    i = n_flag = n_llm = 0
    while i < len(words):
        w, sus, alt = words[i]
        if llm[i][1]:
            runs, j, old = [], i, llm[i][2]
            while j < len(words) and llm[j][1] and llm[j][2] == old:
                if j:
                    _rtl_run(p.add_run(" "))
                r = p.add_run(words[j][0])
                r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
                _rtl_run(r)
                runs.append(r)
                j += 1
            n_llm += len(runs)
            doc.add_comment(runs, text=f"LLM changed — OCR read: {old}", author="LLM", initials="LLM")
            i = j
            continue
        if not sus:
            _rtl_run(p.add_run(("" if i == 0 else " ") + w))
            i += 1
            continue
        runs, j = [], i
        while j < len(words) and words[j][1] and words[j][2] == alt and not llm[j][1]:
            if j:
                _rtl_run(p.add_run(" "))
            r = p.add_run(words[j][0])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            _rtl_run(r)
            runs.append(r)
            j += 1
        n_flag += len(runs)
        doc.add_comment(runs, text=f"{check_name} reads: {alt}", author="OCR", initials="OCR")
        i = j
    return len(words), n_flag, n_llm


def _pairs(main: str, check: str, orig: str | None = None) -> list[tuple[str, str, str | None]]:
    """Same number of lines (two line readers on the same line images; the OCR text before the LLM step too, if
    given) → compare line by line; otherwise (dots reads the whole block) compare the whole block once."""
    a = [x for x in main.split("\n") if x.strip()]
    b = [x for x in check.split("\n") if x.strip()]
    if orig is None:
        return [(x, y, None) for x, y in zip(a, b)] if len(a) == len(b) else [(" ".join(a), " ".join(b), None)]
    c = [x for x in orig.split("\n") if x.strip()]
    if len(a) == len(b) == len(c):
        return list(zip(a, b, c))
    return [(" ".join(a), " ".join(b), " ".join(c))]


def write_docx(path: Path, results: list[PageResult], title: str) -> dict:
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = new_document(title)
    doc.styles["Normal"].font.size = Pt(12)
    rpr = doc.styles["Normal"].element.get_or_add_rPr()  # size of Arabic (complex-script) text: 14 pt
    if rpr.find(qn("w:szCs")) is None:
        el = OxmlElement("w:szCs")
        el.set(qn("w:val"), "28")
        rpr.append(el)
    stats = {"blocks": 0, "words": 0, "words_flagged": 0, "words_llm": 0}
    for k, res in enumerate(results):
        if k:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        llm_ok = bool(res.info.get("llm", {}).get("ok")) and "ocr" in res.texts
        note = doc.add_paragraph().add_run(
            f"Page {k + 1} · {res.info.get('page_type')} · YELLOW = the cross-check reads that word differently "
            "(see the comment) — needs a human check. Hebrew blocks: dots reading, not cross-checked."
            + (" GREEN = changed by the LLM correction (the comment shows what the OCR read) — check it too."
               if llm_ok else ""))
        note.font.size = Pt(9)
        img = res.image
        for bi, b in enumerate(res.blocks):
            cat = b.get("category")
            main, check = res.texts["main"].get(bi, ""), res.texts["check"].get(bi, "")
            if cat in config.TEXT_CATS and main.strip():
                src = res.sources[bi][0]
                orig = res.texts["ocr"].get(bi) if llm_ok and src != "hebrew" else None
                for m, c, o in _pairs(main, check, orig):
                    if src == "hebrew":
                        _set_bidi(doc.add_paragraph(m))
                        continue
                    n, f, g = _marked_paragraph(doc, m, c, "ft3" if src == "dots" else config.CHECK, o)
                    stats["words"] += n
                    stats["words_flagged"] += f
                    stats["words_llm"] += g
                stats["blocks"] += 1
            elif cat == "Table" and b.get("text"):
                _add_markdown(doc, b["text"])
                stats["blocks"] += 1
            elif cat not in config.TEXT_CATS and len(b.get("bbox") or []) == 4:
                x0, y0, x1, y1 = [int(round(v)) for v in b["bbox"]]
                if x1 - x0 > 8 and y1 - y0 > 8:  # stamps, signatures, fingerprints, logos: picture cut from the page
                    buf = io.BytesIO()
                    img.crop((x0, y0, x1, y1)).convert("RGB").save(buf, format="PNG")
                    buf.seek(0)
                    doc.add_picture(buf, width=Cm(min(16.0, (x1 - x0) / img.width * 18)))
                    stats["blocks"] += 1
    doc.save(path)
    return stats


# --- JSON -----------------------------------------------------------------------------------------------------------

def page_json(res: PageResult, page_no: int, stats: dict) -> dict:
    ft3, ft4 = res.readings[config.MAIN], res.readings[config.CHECK]
    return {"page": page_no, **stats, "info": res.info, "image_size": list(res.image.size),
            "blocks": [{"n": bi + 1, "category": b.get("category"), "bbox": b.get("bbox"),
                        "reader": res.sources.get(bi, (None,))[0], "dots_ft3_agreement": res.sources.get(bi, (0, None))[1],
                        "text": res.texts["main"].get(bi), "dots": b.get("text"),
                        **({"text_ocr": res.texts["ocr"].get(bi), "llm_changes": res.llm_changes.get(bi, [])}
                           if "ocr" in res.texts else {}),
                        "lines": [{"bbox": list(L.box), "ft3": ft3[n], "ft4": ft4[n]}
                                  for n, L in enumerate(res.lines) if L.block == bi]}
                       for bi, b in enumerate(res.blocks)]}


# --- everything -----------------------------------------------------------------------------------------------------

def write_all(out_dir: Path, name: str, results: list[PageResult], page_numbers: list[int] | None = None) -> dict:
    """Write every output file of one document → {"files": [paths, zip first], "html": review HTML, "main": text}."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    nums = page_numbers or list(range(1, len(results) + 1))
    n_pages = max(nums) if nums else 0
    htmls, data, boxes = [], [], []
    txt = {k: [] for k in ("main", "ft3", "ft4") + (("ocr",) if any("ocr" in r.texts for r in results) else ())}
    for no, res in zip(nums, results):
        h, st = review_html(res, no, n_pages)
        htmls.append(h)
        data.append(page_json(res, no, st))
        for k in txt:
            txt[k].append(f"===== Page {no} =====\n" + res.page_text(k))
        b = out_dir / f"{name}_p{no}_boxes.jpg"
        draw_boxes(res).save(b, quality=85)
        boxes.append(b)
    files = []
    docx_path = out_dir / f"{name}.docx"
    write_docx(docx_path, results, name)
    files.append(docx_path)
    for k in txt:
        f = out_dir / f"{name}_{k}.txt"
        f.write_text("\n\n".join(txt[k]), encoding="utf-8")
        files.append(f)
    review = out_dir / f"{name}_review.html"
    review.write_text(f"<!doctype html><meta charset=utf-8><title>{html.escape(name)}</title><style>{CSS}</style>"
                      + "".join(htmls), encoding="utf-8")
    js = out_dir / f"{name}.json"
    js.write_text(json.dumps({"document": name, "pages": data}, ensure_ascii=False, indent=1), encoding="utf-8")
    files += [review, js, *boxes]
    z = out_dir / f"{name}_all.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            zf.write(f, f.name)
    return {"files": [z, *files], "html": "".join(htmls), "main": "\n\n".join(txt["main"])}
