"""CPU tests with fake models: text helpers, line segmentation, block → reader choice, all output files.

    python -m pytest -q tests
"""

from __future__ import annotations

import inspect
import json
import zipfile

import numpy as np
import pytest
from PIL import Image, ImageDraw

from arabic_doc_ocr import config
from arabic_doc_ocr.export import write_all
from arabic_doc_ocr.lines import assign_to_blocks, crop_poly, projection_lines
from arabic_doc_ocr.pipeline import Pipeline, load_pages, parse_page_spec, prepare
from arabic_doc_ocr.text import agreement, clean_repeated_substrings, compare_words, fix_arabic, is_hebrew


# --- text -----------------------------------------------------------------------------------------------------------

def test_fix_arabic():
    assert fix_arabic("کتاب یوم ۱۹۷۲") == "كتاب يوم ١٩٧٢"
    assert fix_arabic("و نحن") == "ونحن"
    assert fix_arabic("ذهب و عاد") == "ذهب وعاد"
    assert fix_arabic("نحو بيت") == "نحو بيت"  # و at the end of a word stays


def test_hebrew():
    assert is_hebrew("אישור חתימה 16/5/79")
    assert not is_hebrew("عقد بيع אב")
    assert not is_hebrew("")


def test_compare_words():
    w = compare_words("الفريق الاول باع الارض", "الفريق الأول باع الارص")
    assert [s for _, s, _ in w] == [False, False, False, True]
    assert w[3][2] == "الارص"
    assert agreement("a b c d", "a b c d") == 1.0
    assert agreement("", "x") == 1.0


def test_clean_repeated():
    s = "بداية " + "ab" * 150
    assert len(clean_repeated_substrings(s)) < len(s)
    assert clean_repeated_substrings("short") == "short"


# --- lines ----------------------------------------------------------------------------------------------------------

def _fake_page(n_lines=5, w=1200, h=900) -> Image.Image:
    im = Image.new("RGB", (w, h), (235, 225, 200))
    d = ImageDraw.Draw(im)
    for i in range(n_lines):
        y = 180 + i * 100
        for x in range(100, 1000, 40):
            d.rectangle([x, y, x + 20, y + 30], fill=(30, 30, 120))
    return im


def test_projection_lines():
    boxes = projection_lines(_fake_page(5))
    assert len(boxes) == 5
    ys = [b[1] for b in boxes]
    assert ys == sorted(ys)


def test_assign_and_crop():
    lines = [{"box": (10, 10, 200, 40), "poly": [(10, 10), (200, 10), (200, 40), (10, 40)],
              "baseline": [(10, 35), (200, 35)]},
             {"box": (10, 300, 200, 330), "poly": [], "baseline": []}]
    out = assign_to_blocks(lines, {0: (0, 0, 300, 100), 1: (0, 250, 300, 400), 2: (0, 0, 1000, 1000)})
    assert len(out[0]) == 1 and len(out[1]) == 1 and 2 not in out
    crop = crop_poly(_fake_page(), lines[0])
    assert crop.size == (198, 38)
    assert crop_poly(_fake_page(), lines[0], noi=0).tobytes() == crop.tobytes()  # noi=0 = 2026-10-19 crop


def test_crop_poly_pad():
    # line 10 px thick (y 15–25), ink at y=13 just above it
    im = Image.new("RGB", (60, 40), (250, 250, 250))
    ImageDraw.Draw(im).line([(0, 13), (59, 13)], fill=(0, 0, 0))
    line = {"box": (10, 15, 50, 25), "poly": [(10, 15), (50, 15), (50, 25), (10, 25)]}
    c0, c60 = crop_poly(im, line), crop_poly(im, line, noi=60)  # 60 % of 10 px → 3 px above and below
    assert np.asarray(c0.convert("L"))[13 - 11, 20] > 200      # outside the polygon → paper colour
    assert c60.size == (c0.size[0], c0.size[1] + 6)
    assert np.asarray(c60.convert("L"))[13 - 8, 20] < 50       # ink just above the line is kept
    assert config.LINE_PAD_PCT == 50
    assert inspect.signature(Pipeline.__init__).parameters["line_pad_pct"].default == 50


def test_prepare_and_pages(tmp_path):
    small = Image.new("RGB", (550, 708), "white")
    assert prepare(small).size == (round(550 * 1500 / 708), 1500)
    big = Image.new("RGB", (1032, 1600), "white")
    assert prepare(big).size == (1032, 1600)
    p = tmp_path / "a.jpg"
    big.save(p)
    assert len(load_pages(p)) == 1
    assert parse_page_spec("1-3,5") == [1, 2, 3, 5] and parse_page_spec("") is None


# --- pipeline with fake models --------------------------------------------------------------------------------------

class FakeDots:
    def __init__(self, blocks):
        self.blocks, self.seen = blocks, []

    def parse(self, im):
        self.seen.append(im.size)
        f = im.size[0] / 1200  # boxes come back in the enlarged image's coordinates
        return {"blocks": [{**b, "bbox": [v * f for v in b["bbox"]]} for b in self.blocks], "new_tokens": 10,
                "hit_max_tokens": False, "stopped_loop": False, "layout": "ok"}


class FakeHTR:
    def read(self, crops, adapter=None, progress=None):
        return [("كلمة اولى ثانية" if adapter == "ft3" else "كلمة اولى ثالثة") for _ in crops]

    def read_check(self, ims):
        return [{"text": "a", "check": "a", "words": [("a", False, "")], "agree": 1.0} for _ in ims]


class FakeModels:
    dots_device = htr_device = "cpu"
    offload, loads = "none", {"dots": 1, "htr": 1}

    def __init__(self, blocks):
        self._d, self._h = FakeDots(blocks), FakeHTR()

    def dots(self):
        return self._d

    def htr(self):
        return self._h


BLOCKS = [{"category": "Title", "bbox": [50, 40, 1150, 140], "text": "عقد بيع"},
          {"category": "Text", "bbox": [50, 150, 1150, 680], "text": "كلمة اولى ثانية"},
          {"category": "Picture", "bbox": [800, 700, 1150, 890], "text": ""},
          {"category": "Text", "bbox": [50, 700, 750, 890], "text": "אישור חתימה מס׳ 12"}]


def _pipe(blocks=BLOCKS) -> Pipeline:
    p = Pipeline.__new__(Pipeline)
    p.models = FakeModels(blocks)
    p.line_pad_pct = config.LINE_PAD_PCT
    return p


@pytest.mark.parametrize("page_type", [config.HANDWRITTEN, config.PRINTED])
def test_pipeline_and_outputs(tmp_path, page_type, monkeypatch):
    monkeypatch.setattr("arabic_doc_ocr.pipeline.kraken_page", lambda img: None)  # projection only
    pipe = _pipe()
    res = pipe.process_document([_fake_page()], page_type)
    r = res[0]
    assert pipe.models._d.seen == [(2400, 1800)]  # dots sees the page enlarged to 2400 px
    assert r.blocks[1]["bbox"] == pytest.approx([50, 150, 1150, 680])  # boxes back on the prepared page
    assert r.blocks[3].get("hebrew") and r.sources[3][0] == "hebrew"
    assert 2 not in r.sources  # picture
    assert r.lines and all(L.block != 3 for L in r.lines)  # Hebrew block not given to the line readers
    src = r.sources[1][0]
    assert src == ("dots" if page_type == config.PRINTED else "ft3")
    w = write_all(tmp_path, "doc", res)
    names = {p.name for p in w["files"]}
    assert {"doc_all.zip", "doc.docx", "doc_main.txt", "doc_ft3.txt", "doc_ft4.txt", "doc_review.html", "doc.json",
            "doc_p1_boxes.jpg"} <= names
    js = json.loads((tmp_path / "doc.json").read_text(encoding="utf-8"))
    assert js["pages"][0]["blocks"][3]["reader"] == "hebrew"
    assert "אישור" in (tmp_path / "doc_main.txt").read_text(encoding="utf-8")
    with zipfile.ZipFile(tmp_path / "doc_all.zip") as z:
        assert len(z.namelist()) == len(w["files"]) - 1
    from docx import Document

    doc = Document(tmp_path / "doc.docx")
    assert any("كلمة" in p.text for p in doc.paragraphs)


def test_fallback_when_no_text_block(monkeypatch):
    monkeypatch.setattr("arabic_doc_ocr.pipeline.kraken_page", lambda img: None)
    pipe = _pipe([{"category": "Picture", "bbox": [0, 0, 1200, 900], "text": ""}])
    r = pipe.process(_fake_page(), config.HANDWRITTEN)
    assert r.blocks[0].get("fallback") and len(r.lines) == 5
    assert r.info["notes"] and "FALLBACK" in r.info["notes"][0]


def test_oom_retry(monkeypatch):
    monkeypatch.setattr("arabic_doc_ocr.pipeline.kraken_page", lambda img: None)
    pipe = _pipe()
    real = pipe.models._d.parse

    def parse(im):
        if im.size[0] > 1200:
            raise RuntimeError("CUDA out of memory")
        return real(im)

    pipe.models._d.parse = parse
    r = pipe.layout(_fake_page())
    assert r.info["dots_scale"] == 1.0 and "out of GPU memory" in r.info["notes"][0]


def test_vendored_dots_postprocess():
    """The vendored dots utils turn a layout JSON into blocks scaled to the original image."""
    from arabic_doc_ocr.dots_model import _dots_utils, layout_loop

    _, _, layout_utils = _dots_utils()
    origin = Image.new("RGB", (1000, 1400))
    resp = json.dumps([{"bbox": [10, 20, 500, 80], "category": "Text", "text": "نص"}])
    cells, filtered = layout_utils.post_process_output(resp, config.DOTS_PROMPT, origin, origin)
    assert not filtered and cells[0]["text"] == "نص"
    loop = "".join('{"bbox": [1, 2, 3, 4], "category": "Text", "text": "abcdefg"}, ' for _ in range(30))
    assert layout_loop(loop) and not layout_loop(loop[:200])


def test_projection_on_real_scan_shapes():
    """Ruled paper: long horizontal rules must not become lines."""
    im = _fake_page(3)
    d = ImageDraw.Draw(im)
    for y in (160, 260, 360, 460):
        d.line([0, y, 1199, y], fill=(30, 30, 120), width=2)
    assert len(projection_lines(im)) == 3
    assert np.asarray(im).shape == (900, 1200, 3)


# --- LLM correction (CPU, fake LLM — the real one runs in its own venv, scripts/check_llm.py) ----------------------

class FakeLLM:
    """Prefers any sentence containing "ثالثة"; chat() returns a fixed rewrite."""

    def __init__(self, reply=""):
        self.reply = reply

    def logprob(self, text):
        return 1.0 if "ثالثة" in text else 0.0

    def chat(self, user):
        return self.reply


def test_llm_chon_picks_the_more_probable_reading():
    from arabic_doc_ocr.llm_fix import fix_chon

    t, ch = fix_chon(FakeLLM(), "كلمة اولى ثانية", "كلمة اولى ثالثة")
    assert t == "كلمة اولى ثالثة" and ch[0]["old"] == "ثانية" and ch[0]["new"] == "ثالثة"
    t, ch = fix_chon(FakeLLM(), "رقم 12 ثانية", "رقم 13 ثالثة")  # spans with digits are never touched
    assert t == "رقم 12 ثانية" and ch == []


def test_llm_sua_accepts_only_small_word_edits():
    from arabic_doc_ocr.llm_fix import fix_sua

    # 1-letter fixes accepted; a rewritten word (too many letters) and a changed number rejected
    t, ch = fix_sua(FakeLLM("الفريق الأول اشترى الارض سنة 1973"), "الفريق الاول باع الارص سنة 1972", "")
    assert t == "الفريق الأول باع الارض سنة 1972"
    assert sorted(c["new"] for c in ch if not c.get("rejected")) == ["الأول", "الارض"]
    assert sum(1 for c in ch if c.get("rejected")) == 2


def test_llm_not_installed_reason(monkeypatch, tmp_path):
    from arabic_doc_ocr import llm

    monkeypatch.setattr(config, "LLM_VENV", tmp_path / "nothing")
    assert llm.llm_python() is None and "install_llm.sh" in llm.why_not()


def test_llm_apply_and_outputs(tmp_path, monkeypatch):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    from arabic_doc_ocr import llm

    monkeypatch.setattr("arabic_doc_ocr.pipeline.kraken_page", lambda img: None)
    res = _pipe().process_document([_fake_page()], config.HANDWRITTEN)
    us = llm.units(res)
    assert us and {u["key"] for u in us} == {f"0:{bi}" for bi, (s, _) in res[0].sources.items()
                                              if s in ("dots", "ft3") and res[0].texts["main"][bi].strip()}
    r = {"texts": {u["key"]: u["main"].replace("ثانية", "ثالثة") for u in us},
         "changes": {u["key"]: [{"old": "ثانية", "new": "ثالثة", "step": "chon"},
                                {"old": "اولى", "new": "اول", "step": "sua", "rejected": True}] for u in us}}
    assert llm.apply(res, r) == len(us)  # rejected proposals are not counted
    res[0].info["llm"] = {"ok": True, "model": "fake", "method": "chon_sua"}
    bi = int(next(u for u in us if "ثانية" in u["main"])["key"].split(":")[1])
    assert "ثالثة" in res[0].texts["main"][bi] and "ثانية" in res[0].texts["ocr"][bi]
    w = write_all(tmp_path, "doc", res)
    assert "doc_ocr.txt" in {p.name for p in w["files"]}
    assert "ثانية" in (tmp_path / "doc_ocr.txt").read_text(encoding="utf-8")
    assert "ثانية" not in (tmp_path / "doc_main.txt").read_text(encoding="utf-8")
    js = json.loads((tmp_path / "doc.json").read_text(encoding="utf-8"))
    b = js["pages"][0]["blocks"][bi]
    assert "ثانية" in b["text_ocr"] and b["llm_changes"][0]["new"] == "ثالثة"
    assert "After the LLM correction" in (tmp_path / "doc_review.html").read_text(encoding="utf-8")
    doc = Document(tmp_path / "doc.docx")
    green = [x.text for p in doc.paragraphs for x in p.runs if x.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN]
    assert "ثالثة" in green and "ثانية" not in green
